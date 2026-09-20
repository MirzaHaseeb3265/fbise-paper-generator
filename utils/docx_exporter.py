from __future__ import annotations
from io import BytesIO
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def _setup(doc):
    sec=doc.sections[0]; sec.page_width=Mm(210); sec.page_height=Mm(297)
    sec.top_margin=Mm(10); sec.bottom_margin=Mm(10); sec.left_margin=Mm(12); sec.right_margin=Mm(12)
    style=doc.styles['Normal']; style.font.name='Arial'; style.font.size=Pt(10)
    pf=style.paragraph_format; pf.space_after=Pt(1); pf.space_before=Pt(0); pf.line_spacing=1.0

def _p(doc,text="",bold=False,center=False,size=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(1)
    r=p.add_run(text); r.bold=bold
    if size:r.font.size=Pt(size)
    if center:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    return p

def build_docx(paper,s):
    doc=Document(); _setup(doc)
    title="Physics HSC-I - FLP" if s['paper_type']=="FLP Paper" and s['level']=="HSSC-I" else ("Physics HSC-II - FLP" if s['paper_type']=="FLP Paper" else f"FSC {'PART-I' if s['level']=='HSSC-I' else 'PART-II'} (FEDERAL)")
    _p(doc,title,True,True,12); _p(doc,"Topics: "+", ".join(s['chapters']),False,True,9)
    if s['paper_type']=="FLP Paper": _p(doc,"SECTION-A (OBJECTIVE)",True,True); _p(doc,"Q No 1: MULTIPLE CHOICE QUESTIONS",True)
    else:_p(doc,f"Encircle the right answer ({len(paper['mcqs'])*s['mcq_marks']} Marks)",True)
    for i,q in enumerate(paper['mcqs'],1):
        _p(doc,f"{i}. {q['question']}")
        o=q['options']; _p(doc,f"A. {o['A']}    B. {o['B']}    C. {o['C']}    D. {o['D']}")
    _p(doc,"SECTION-B (SHORT QUESTIONS)" if s['paper_type']=="FLP Paper" else "Short Questions",True,True)
    _p(doc,f"Attempt {'any '+str(s.get('short_attempt',len(paper['short_questions']))) if s.get('short_attempt',len(paper['short_questions']))<len(paper['short_questions']) else 'all'} questions. ({s['short_marks']} marks each)")
    for i,q in enumerate(paper['short_questions'],1):
        _p(doc,f"{i}. {q['question']}")
        if q.get('or_question'): _p(doc,"OR",True,True); _p(doc,q['or_question'])
    _p(doc,"SECTION-C (LONG QUESTIONS)" if s['paper_type']=="FLP Paper" else "Long Questions",True,True)
    _p(doc,f"Attempt {'any '+str(s.get('long_attempt',len(paper['long_questions']))) if s.get('long_attempt',len(paper['long_questions']))<len(paper['long_questions']) else 'all'} questions. ({s['long_marks']} marks each)")
    for i,q in enumerate(paper['long_questions'],1):
        _p(doc,f"{i}. {q['question']}")
        if q.get('or_question'): _p(doc,"OR",True,True); _p(doc,q['or_question'])
    b=BytesIO(); doc.save(b); return b.getvalue()

def build_answer_key_docx(paper,s):
    doc=Document(); _setup(doc); _p(doc,"Teacher Answer Key",True,True,13)
    _p(doc,"MCQs",True)
    for i,q in enumerate(paper['mcqs'],1): _p(doc,f"{i}. {q.get('correct_answer','')} — {q.get('answer_explanation','')} [p. {', '.join(map(str,q.get('source_pages',[])))}]")
    _p(doc,"Short Questions",True)
    for i,q in enumerate(paper['short_questions'],1): _p(doc,f"{i}. " + "; ".join(q.get('answer_key_points',[])))
    _p(doc,"Long Questions",True)
    for i,q in enumerate(paper['long_questions'],1): _p(doc,f"{i}. " + "; ".join(q.get('marking_outline',[])))
    b=BytesIO(); doc.save(b); return b.getvalue()
